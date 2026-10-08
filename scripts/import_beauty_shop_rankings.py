#!/usr/bin/env python3
"""Normalize Shopee TW Beauty monthly shop-ranking exports into static JSON.

Requires openpyxl. Usage:
  python scripts/import_beauty_shop_rankings.py <xlsx-directory> [output-json]
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


FILE_PREFIX = "[ExportReport]BrandPortal_MarketIntelligence_Shop+Ranking_"
OUTSIDE_TOP10_HEADER = "My shop outside the Top-10 ranks"
SHEETS = {
    "sales": ("By Sales", "Sales(NT$)", "Sales(USD)"),
    "units": ("By Units Sold", "Units Sold", None),
    "views": ("By Product Views", "Product View", None),
    "transactionIndex": ("By Transaction Index", "Transaction Index", None),
}


def number_or_none(value: object) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        parsed = float(text)
    except ValueError as exc:
        raise ValueError(f"Expected a number or blank, got {value!r}") from exc
    return int(parsed) if parsed.is_integer() else parsed


def normalize_shop_url(value: object) -> str:
    parsed = urlsplit(str(value).strip())
    if not parsed.netloc:
        raise ValueError(f"Shop URL has no host: {value!r}")
    path = parsed.path.rstrip("/")
    return f"https://{parsed.netloc.lower()}{path}"


def month_from_file(path: Path) -> str:
    match = re.fullmatch(r"(\d{2})_2026\.xlsx", path.name[len(FILE_PREFIX) :])
    if not match:
        raise ValueError(f"Unexpected ranking filename: {path.name}")
    return f"2026-{match.group(1)}"


def rows_for_month(path: Path, month: str) -> dict[str, list[dict[str, object]]]:
    # These exports contain an incomplete worksheet dimension in the XLSX XML;
    # normal mode is required to read all columns and rows reliably.
    workbook = load_workbook(path, data_only=True, read_only=False)
    definitions = workbook["Definitions"]
    if str(definitions["B3"].value).strip().upper() != "TW":
        raise ValueError(f"{path.name}: expected TW region")
    if str(definitions["B4"].value).strip().lower() != "beauty":
        raise ValueError(f"{path.name}: expected Beauty category")
    period = str(definitions["B5"].value or "")
    period_start = period.split("~", maxsplit=1)[0].strip()
    parsed_start = datetime.strptime(period_start, "%d/%m/%Y")
    if parsed_start.strftime("%Y-%m") != month:
        raise ValueError(f"{path.name}: month {month} does not match source period {period!r}")

    result: dict[str, list[dict[str, object]]] = {}
    for metric, (sheet_name, value_header, secondary_header) in SHEETS.items():
        worksheet = workbook[sheet_name]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in worksheet[1]]
        indices = {header: index for index, header in enumerate(headers) if header}
        required = {"Ranking", "Change of Rank", "Shop Name", "Shop Link", "Is it My Shop", value_header}
        if secondary_header:
            required.add(secondary_header)
        missing = required - indices.keys()
        if missing:
            raise ValueError(f"{path.name}/{sheet_name}: missing headers {sorted(missing)}")

        metric_rows: list[dict[str, object]] = []
        seen_urls: set[str] = set()
        ranked_count = 0
        outside_top10_section = False
        outside_top10_position = 0
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if any(str(value).strip() == OUTSIDE_TOP10_HEADER for value in values if value is not None):
                outside_top10_section = True
                continue

            raw_url = values[indices["Shop Link"]]
            if not raw_url:
                continue
            shop_url = normalize_shop_url(raw_url)
            if shop_url in seen_urls:
                raise ValueError(f"{path.name}/{sheet_name}: duplicate shop URL {shop_url}")
            seen_urls.add(shop_url)

            raw_rank = values[indices["Ranking"]]
            rank_text = str(raw_rank).strip()
            if rank_text.isdigit():
                rank: int | None = int(rank_text)
                rank_status = "ranked"
                rank_basis = "source"
                ranked_count += 1
            elif rank_text == "-":
                if outside_top10_section:
                    outside_top10_position += 1
                    rank = 10 + outside_top10_position
                    rank_status = "ranked"
                    rank_basis = "orderedList"
                else:
                    rank = None
                    rank_status = "outsideTop10"
                    rank_basis = None
            else:
                raise ValueError(f"{path.name}/{sheet_name}: unsupported rank {raw_rank!r}")

            record: dict[str, object] = {
                "month": month,
                "shopName": str(values[indices["Shop Name"]] or "").strip(),
                "shopUrl": shop_url,
                "rank": rank,
                "rankStatus": rank_status,
                "rankBasis": rank_basis,
                "sourceRank": rank_text,
                "rankChange": (
                    str(values[indices["Change of Rank"]]).strip()
                    if values[indices["Change of Rank"]] is not None
                    else None
                ),
                "value": number_or_none(values[indices[value_header]]),
                "isMyShop": str(values[indices["Is it My Shop"]] or "").strip().upper() == "Y",
            }
            if secondary_header:
                record["valueUsd"] = number_or_none(values[indices[secondary_header]])
            metric_rows.append(record)

        if ranked_count != 10:
            raise ValueError(f"{path.name}/{sheet_name}: expected 10 ranked shops, found {ranked_count}")
        if outside_top10_position < 10:
            raise ValueError(
                f"{path.name}/{sheet_name}: expected at least 10 ordered rows after "
                f"{OUTSIDE_TOP10_HEADER!r}, found {outside_top10_position}"
            )
        result[metric] = metric_rows

    workbook.close()
    return result


def main() -> int:
    if len(sys.argv) not in (2, 3):
        print(__doc__.strip(), file=sys.stderr)
        return 2

    source_dir = Path(sys.argv[1]).expanduser().resolve()
    output_path = (
        Path(sys.argv[2]).expanduser().resolve()
        if len(sys.argv) == 3
        else Path(__file__).resolve().parents[1] / "public" / "beauty_shop_rankings_2026.json"
    )
    files = sorted(
        path
        for path in source_dir.iterdir()
        if path.is_file() and path.name.startswith(FILE_PREFIX) and path.name.endswith("_2026.xlsx")
    )
    if not files:
        raise FileNotFoundError(f"No 2026 ranking exports found in {source_dir}")

    months = [month_from_file(path) for path in files]
    if len(months) != len(set(months)):
        raise ValueError("Duplicate month exports found")
    if months != sorted(months):
        raise ValueError("Ranking exports are not in chronological month order")

    normalized: dict[str, list[dict[str, object]]] = {metric: [] for metric in SHEETS}
    for path, month in zip(files, months, strict=True):
        monthly_rows = rows_for_month(path, month)
        for metric, rows in monthly_rows.items():
            normalized[metric].extend(rows)

    payload = {
        "region": "TW",
        "category": "Beauty",
        "timezone": "Asia/Taipei",
        "rankingLimit": 10,
        "months": months,
        "latestMonth": months[-1],
        "updatedAt": date.today().isoformat(),
        "sourceFiles": [path.name for path in files],
        "metricLabels": {
            "sales": "Sales",
            "units": "Units Sold",
            "views": "Product Views",
            "transactionIndex": "Transaction Index",
        },
        "metrics": normalized,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    counts = {metric: len(rows) for metric, rows in normalized.items()}
    print(json.dumps({"output": str(output_path), "months": months, "rowCounts": counts}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
